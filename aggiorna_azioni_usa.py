#!/usr/bin/env python3
"""Aggiorna il foglio "Azioni USA" con le azioni USA oltre 1 miliardo di capitalizzazione.

Colonne scritte: capitalizzazione, prezzo, media mobile a 52 settimane (252 chiusure
giornaliere), RSI a 1 ora, 1 giorno e 1 settimana con stato (ipercomprato/ipervenduto).

Uso:
    pip install yfinance pandas openpyxl
    python aggiorna_azioni_usa.py --workbook simulatore_investimento_corretto.xlsx

Chiudi il file Excel prima di eseguire lo script.
Fonte dati: Yahoo Finance tramite la libreria yfinance (dati ritardati, uso personale).
Se lo screener di Yahoo non risponde, crea un file tickers.csv con una colonna "symbol".
"""
import argparse
import datetime as dt
import sys

import numpy as np
import pandas as pd

RSI_PERIOD = 14
OVERBOUGHT = 70
OVERSOLD = 30
SMA_DAYS = 252          # circa 52 settimane di borsa
MIN_CAP = 1_000_000_000
CHUNK = 100             # ticker per richiesta
SHEET = "Azioni USA"
FIRST_ROW = 6           # prima riga dati nel foglio


def rsi(close, n=RSI_PERIOD):
    """RSI di Wilder sull'ultima barra (media iniziale semplice, poi smoothing di Wilder).
    Restituisce NaN se i dati non bastano."""
    close = pd.Series(close).dropna().to_numpy(dtype=float)
    if len(close) < n + 1:
        return np.nan
    delta = np.diff(close)
    gain = np.where(delta > 0, delta, 0.0)
    loss = np.where(delta < 0, -delta, 0.0)
    avg_gain = gain[:n].mean()
    avg_loss = loss[:n].mean()
    for g, l in zip(gain[n:], loss[n:]):
        avg_gain = (avg_gain * (n - 1) + g) / n
        avg_loss = (avg_loss * (n - 1) + l) / n
    if avg_loss == 0:
        return 100.0 if avg_gain > 0 else 50.0
    return 100 - 100 / (1 + avg_gain / avg_loss)


def status(value):
    if value is None or np.isnan(value):
        return "n.d."
    if value >= OVERBOUGHT:
        return "Ipercomprato"
    if value <= OVERSOLD:
        return "Ipervenduto"
    return "Neutro"


def sma(close, days=SMA_DAYS):
    close = pd.Series(close).dropna()
    if len(close) < int(days * 0.8):
        return np.nan
    return float(close.tail(days).mean())


def get_universe():
    """Elenco azioni USA con capitalizzazione > 1 miliardo, ordinate per capitalizzazione."""
    import yfinance as yf
    try:
        from yfinance import EquityQuery
        query = EquityQuery("and", [
            EquityQuery("eq", ["region", "us"]),
            EquityQuery("is-in", ["exchange", "NMS", "NYQ", "NGM", "ASE"]),
            EquityQuery("gt", ["intradaymarketcap", MIN_CAP]),
        ])
        rows, offset = [], 0
        while True:
            res = yf.screen(query, offset=offset, size=250,
                            sortField="intradaymarketcap", sortAsc=False)
            quotes = res.get("quotes", [])
            if not quotes:
                break
            for q in quotes:
                rows.append({
                    "symbol": q.get("symbol"),
                    "name": q.get("shortName") or q.get("longName") or "",
                    "marketCap": q.get("marketCap"),
                    "price": q.get("regularMarketPrice"),
                })
            offset += len(quotes)
            if offset >= res.get("total", 0):
                break
        df = pd.DataFrame(rows).dropna(subset=["symbol"])
        if not df.empty:
            return df[df["marketCap"].fillna(0) >= MIN_CAP]
    except Exception as exc:  # screener non disponibile
        print(f"Screener Yahoo non disponibile ({exc}); provo tickers.csv", file=sys.stderr)
    try:
        tick = pd.read_csv("tickers.csv")
        tick = tick.rename(columns={tick.columns[0]: "symbol"})
        tick["name"] = ""
        tick["marketCap"] = np.nan
        tick["price"] = np.nan
        return tick[["symbol", "name", "marketCap", "price"]]
    except Exception:
        sys.exit("Nessun elenco di titoli: serve lo screener Yahoo oppure un file tickers.csv")


def fetch_closes(tickers, interval, period):
    """Scarica le chiusure per un elenco di ticker. Restituisce {ticker: Series}."""
    import yfinance as yf
    out = {}
    for i in range(0, len(tickers), CHUNK):
        part = tickers[i:i + CHUNK]
        try:
            data = yf.download(part, period=period, interval=interval, group_by="ticker",
                               auto_adjust=True, progress=False, threads=True)
        except Exception as exc:
            print(f"Errore download {interval} {part[0]}...: {exc}", file=sys.stderr)
            continue
        for t in part:
            try:
                col = data[t]["Close"] if len(part) > 1 else data["Close"]
                s = col.dropna()
                if not s.empty:
                    out[t] = s
            except Exception:
                pass
        print(f"  {interval}: {min(i + CHUNK, len(tickers))}/{len(tickers)}", file=sys.stderr)
    return out


def build_table(universe, daily, weekly, hourly):
    rows = []
    for _, u in universe.iterrows():
        t = u["symbol"]
        d, w, h = daily.get(t), weekly.get(t), hourly.get(t)
        price = u["price"]
        if (price is None or pd.isna(price)) and d is not None:
            price = float(d.iloc[-1])
        media = sma(d) if d is not None else np.nan
        r_h = rsi(h) if h is not None else np.nan
        r_d = rsi(d) if d is not None else np.nan
        r_w = rsi(w) if w is not None else np.nan
        rows.append({
            "Ticker": t, "Società": u["name"],
            "Capitalizzazione (USD)": u["marketCap"], "Prezzo (USD)": price,
            "Media 52 sett. (USD)": media,
            "Prezzo vs media": (price / media - 1) if media and not np.isnan(media) and price else np.nan,
            "RSI 1h": r_h, "Stato 1h": status(r_h),
            "RSI 1g": r_d, "Stato 1g": status(r_d),
            "RSI 1s": r_w, "Stato 1s": status(r_w),
        })
    return pd.DataFrame(rows)


FORMATS = {1: None, 2: None, 3: "#,##0", 4: "#,##0.00", 5: "#,##0.00", 6: "0.0%",
           7: "0.0", 9: "0.0", 11: "0.0"}


def write_workbook(path, table, stamp):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[SHEET]
    ws["B3"] = stamp
    if ws.max_row >= FIRST_ROW:
        ws.delete_rows(FIRST_ROW, ws.max_row - FIRST_ROW + 1)
    for i, rec in enumerate(table.itertuples(index=False)):
        for j, v in enumerate(rec, start=1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            c = ws.cell(FIRST_ROW + i, j, v)
            if FORMATS.get(j):
                c.number_format = FORMATS[j]
    ws.auto_filter.ref = f"A{FIRST_ROW - 1}:L{FIRST_ROW - 1 + len(table)}"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="simulatore_investimento_corretto.xlsx")
    ap.add_argument("--csv", default=None, help="salva anche un CSV")
    ap.add_argument("--limit", type=int, default=None, help="solo i primi N titoli (prove)")
    args = ap.parse_args()

    universe = get_universe().sort_values("marketCap", ascending=False)
    if args.limit:
        universe = universe.head(args.limit)
    tickers = universe["symbol"].tolist()
    print(f"{len(tickers)} titoli", file=sys.stderr)
    daily = fetch_closes(tickers, "1d", "2y")
    weekly = fetch_closes(tickers, "1wk", "5y")
    hourly = fetch_closes(tickers, "1h", "60d")
    table = build_table(universe, daily, weekly, hourly)
    stamp = dt.datetime.now().strftime("%d/%m/%Y %H:%M")
    if args.csv:
        table.to_csv(args.csv, index=False)
    write_workbook(args.workbook, table, stamp)
    print(f"Aggiornato il foglio '{SHEET}' alle {stamp}", file=sys.stderr)


if __name__ == "__main__":
    main()
